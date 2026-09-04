#!/usr/bin/env python3
"""Собирает строки учёта времени из событий календаря.

Вход:  timesheet/events/<ГГГГ-ММ-ДД>.json  — события дня (снимаются со скриншота)
       timesheet/rules.json                — правила: проект, тип работ, исключения
Выход: timesheet/out/<ГГГГ-ММ-ДД>_<лист>.xlsx — файл Excel
       timesheet/out/<ГГГГ-ММ-ДД>_<лист>.tsv  — для вставки прямо в Google Sheets

Колонки повторяют лист «Оксана»: дата | к-во часов | проект | тип работ | комментарий

    python3 timesheet/build_timesheet.py 2026-08-10
"""

import json
import os
import sys
import urllib.request
from datetime import datetime
from pathlib import Path

BASE = Path(__file__).resolve().parent
COLUMNS = ["дата", "к-во часов", "проект", "тип работ", "комментарий"]


def match_rule(title, rules, key):
    """Первое подошедшее правило.

    `match` срабатывает по любой из подстрок, `all` — только если в названии
    есть все перечисленные сразу. Второе нужно там, где решает сочетание слов:
    «Росинка» вместе с «разработкой» — это отдельный проект, а порознь нет.
    """
    low = title.lower()
    for rule in rules:
        required = rule.get("all")
        if required and all(p.lower() in low for p in required):
            return rule[key], " + ".join(required)
        for pattern in rule.get("match", []):
            if pattern.lower() in low:
                return rule[key], pattern
    return None, None


def hours_between(start, end, step):
    fmt = "%H:%M"
    delta = datetime.strptime(end, fmt) - datetime.strptime(start, fmt)
    hours = delta.total_seconds() / 3600
    return round(round(hours / step) * step, 2)


def build_rows(day, rules):
    rows, skipped, warnings = [], [], []

    for event in day["events"]:
        title = event["title"]

        if event.get("all_day"):
            skipped.append((title, "событие на весь день"))
            continue

        # в календаре отклонённые встречи показаны зачёркнутыми
        if event.get("declined"):
            skipped.append((title, "встреча отклонена"))
            continue

        excluded = next(
            (word for word in rules["exclude"] if word.lower() in title.lower()), None
        )
        if excluded:
            skipped.append((title, f"исключено правилом «{excluded}»"))
            continue

        hours = hours_between(event["start"], event["end"], rules["round_hours_to"])
        if hours <= 0:
            skipped.append((title, "нулевая длительность"))
            continue

        project, _ = match_rule(title, rules["projects"], "project")
        if project is None:
            project = rules["default_project"]
            warnings.append(f"«{title}»: проект не распознан → {project}")

        work_type, _ = match_rule(title, rules["work_types"], "work_type")
        if work_type is None:
            work_type = rules["default_work_type"]
            warnings.append(f"«{title}»: тип работ не распознан → {work_type}")

        if event.get("check"):
            warnings.append(f"«{title}»: {event['check']}")

        if event.get("uncertain"):
            warnings.append(
                f"«{title}»: время снято со скриншота приблизительно "
                f"({event['start']}–{event['end']}) — {event.get('note', '')}".strip()
            )

        rows.append(
            {
                "hours": hours,
                "project": project,
                "work_type": work_type,
                # исходное название события — для проверки; в лист попадает,
                # только если в правилах включён emit_comment
                "comment": title if rules.get("emit_comment") else "",
                "source": title,
            }
        )

    if rules["merge_same_project_and_type"]:
        rows = merge(rows)

    return rows, skipped, warnings


def merge(rows):
    merged = {}
    for row in rows:
        key = (row["project"], row["work_type"])
        if key in merged:
            merged[key]["hours"] = round(merged[key]["hours"] + row["hours"], 2)
            for field in ("comment", "source"):
                parts = merged[key][field].split("; ")
                if row[field] and row[field] not in parts:
                    merged[key][field] = "; ".join(filter(None, parts + [row[field]]))
        else:
            merged[key] = dict(row)
    return list(merged.values())


def ru_date(iso):
    return datetime.strptime(iso, "%Y-%m-%d").strftime("%d.%m.%y")


def ru_number(value):
    text = f"{value:.2f}".rstrip("0").rstrip(".")
    return text.replace(".", ",")


def write_tsv(path, iso_date, rows):
    lines = ["\t".join(COLUMNS)]
    for row in rows:
        lines.append(
            "\t".join(
                [
                    ru_date(iso_date),
                    ru_number(row["hours"]),
                    row["project"],
                    row["work_type"],
                    row["comment"],
                ]
            )
        )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_paste_tsv(path, rows):
    """Только колонки B:E — в листе даты за день уже проставлены заранее."""
    lines = [
        "\t".join([ru_number(row["hours"]), row["project"], row["work_type"], row["comment"]])
        for row in rows
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_xlsx(path, sheet_name, iso_date, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name

    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for row in rows:
        sheet.append(
            [
                ru_date(iso_date),
                row["hours"],
                row["project"],
                row["work_type"],
                row["comment"],
            ]
        )

    for index, width in enumerate([10, 12, 28, 62, 46], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A2"
    workbook.save(path)


def push(rules, iso_date, rows):
    """Отправляет строки в Apps Script, привязанный к таблице."""
    url = rules.get("webhook_url")
    secret = os.environ.get("TIMESHEET_SECRET")

    if not url:
        sys.exit("Нет webhook_url в rules.json — сначала разверните timesheet/appscript/Code.gs")
    if not secret:
        sys.exit("Нет переменной окружения TIMESHEET_SECRET")

    payload = {
        "secret": secret,
        "date": ru_date(iso_date),
        "rows": [
            {
                "hours": row["hours"],
                "project": row["project"],
                "work_type": row["work_type"],
                "comment": row["comment"],
            }
            for row in rows
        ],
    }

    request = urllib.request.Request(
        url,
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
    )
    with urllib.request.urlopen(request, timeout=30) as response:
        result = json.loads(response.read().decode("utf-8"))

    if result.get("status") != "ok":
        sys.exit(f"Таблица вернула ошибку: {result.get('message')}")

    replaced = result.get("replaced", 0)
    note = f", заменено прежних строк за эту дату: {replaced}" if replaced else ""
    print(f"\nЗаписано в лист «{rules['sheet']}»: {result.get('added')} строк{note}")


def main():
    args = sys.argv[1:]
    should_push = "--push" in args
    args = [arg for arg in args if arg != "--push"]

    if len(args) != 1:
        sys.exit("Использование: python3 timesheet/build_timesheet.py ГГГГ-ММ-ДД [--push]")

    iso_date = args[0]
    events_path = BASE / "events" / f"{iso_date}.json"
    if not events_path.exists():
        sys.exit(f"Нет файла событий: {events_path}")

    rules = json.loads((BASE / "rules.json").read_text(encoding="utf-8"))
    day = json.loads(events_path.read_text(encoding="utf-8"))

    rows, skipped, warnings = build_rows(day, rules)
    rows.sort(key=lambda row: (row["project"], row["work_type"]))

    out_dir = BASE / "out"
    out_dir.mkdir(exist_ok=True)
    stem = f"{iso_date}_{rules['sheet']}"
    write_tsv(out_dir / f"{stem}.tsv", iso_date, rows)
    write_paste_tsv(out_dir / f"{stem}_вставка_B-E.tsv", rows)
    write_xlsx(out_dir / f"{stem}.xlsx", rules["sheet"], iso_date, rows)

    total = round(sum(row["hours"] for row in rows), 2)
    print(f"{ru_date(iso_date)} — лист «{rules['sheet']}», строк: {len(rows)}, часов: {ru_number(total)}\n")
    for row in rows:
        print(f"  {ru_number(row['hours']):>5}  {row['project']:<26} {row['work_type'].strip()}")
        print(f"         ← {row['source']}")

    if skipped:
        print("\nНе попало в учёт:")
        for title, reason in skipped:
            print(f"  — {title}: {reason}")

    if warnings:
        print("\nПроверьте вручную:")
        for warning in warnings:
            print(f"  ! {warning}")

    print(f"\nФайлы: {out_dir / (stem + '.xlsx')}\n       {out_dir / (stem + '.tsv')}")

    if should_push:
        push(rules, iso_date, rows)


if __name__ == "__main__":
    main()
